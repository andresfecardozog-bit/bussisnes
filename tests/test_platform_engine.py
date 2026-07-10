"""Tests del motor generico: cero perdida, grano, KPIs, y la regresion
PRE CORTE (gate duro de Fase 1: KPIs identicos al pipeline legado)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from app.platform.engine import (
    GranoNoResueltoError,
    run_profile,
    run_profile_multi,
)
from app.platform.profile import MatchProfile

FIXTURES = Path(__file__).parent / "fixtures"
PROFILES_DIR = Path(__file__).resolve().parents[1] / "profiles"
PRE_CORTE_XLSX = FIXTURES / "PRE_CORTE_muestra.xlsx"
FLASH_CSV = FIXTURES / "FLASH_muestra.csv"
CEN_XLSX = FIXTURES / "cen" / "Acumulado CEN P7 2026.xlsx"
CEN_JUNIO = FIXTURES / "cen" / "cen_junio_muestra.xlsx"
SAP_MUESTRA = FIXTURES / "cen" / "sap_junio_muestra.xlsx"

FECHA = date(2026, 2, 14)


def _pre_corte_profile() -> MatchProfile:
    raw = (PROFILES_DIR / "pre_corte_v1.json").read_text(encoding="utf-8")
    return MatchProfile.from_json(raw)


def _cen_profile() -> MatchProfile:
    raw = (PROFILES_DIR / "cen_vs_sap_v1_borrador.json").read_text(encoding="utf-8")
    return MatchProfile.from_json(raw)


# ---------------------------------------------------------------------------
# Regresion PRE CORTE: el profile reproduce el pipeline legado
# ---------------------------------------------------------------------------

def test_regresion_pre_corte_kpis_identicos_al_legado(isolated_db_with_catalog):
    from app.core.matcher import run_full_pipeline

    legado = run_full_pipeline(PRE_CORTE_XLSX, FLASH_CSV, fecha_override=FECHA)

    result = run_profile(
        _pre_corte_profile(),
        left_path=PRE_CORTE_XLSX,
        right_path=FLASH_CSV,
        parameters={"fecha_produccion": FECHA},
    )

    # mismas particiones
    assert result.summary()["matched"] == legado.summary()["matched"]
    assert result.summary()["solo_left"] == legado.summary()["solo_pre_corte"]
    assert result.summary()["solo_right"] == legado.summary()["solo_flash"]
    assert result.summary()["no_cruzados"] == legado.summary()["no_cruzados"]

    # KPIs identicos: totales y cumplimiento global
    plan_legado = float(legado.matched["notificado_unidades"].sum())
    real_legado = float(legado.matched["real_unidades_flash"].sum())
    assert result.kpis["plan_total_unidades"] == plan_legado
    assert result.kpis["real_total_unidades"] == real_legado
    esperado = round(real_legado / plan_legado * 100.0, 2)
    assert result.kpis["cumplimiento_global_pct"] == esperado

    # cumplimiento por material identico fila a fila
    gen = result.matched.set_index("_k0")["cumplimiento_pct"].sort_index()
    leg = legado.matched.set_index("material")["cumplimiento_pct"].sort_index()
    assert list(gen.round(2)) == list(leg.round(2))


def test_regresion_pre_corte_no_cruzados_motivos(isolated_db_with_catalog):
    result = run_profile(
        _pre_corte_profile(),
        left_path=PRE_CORTE_XLSX,
        right_path=FLASH_CSV,
        parameters={"fecha_produccion": FECHA},
    )
    if not result.no_cruzados.empty:
        motivos = set(result.no_cruzados["motivo"])
        assert motivos <= {
            "material notificado sin venta en FLASH",
            "material vendido sin notificacion previa",
        }


def test_parametro_requerido_faltante(isolated_db_with_catalog):
    with pytest.raises(ValueError, match="faltantes"):
        run_profile(
            _pre_corte_profile(),
            left_path=PRE_CORTE_XLSX,
            right_path=FLASH_CSV,
            parameters={},
        )


# ---------------------------------------------------------------------------
# CEN vs SAP con el borrador a mano (fixtures recortados)
# ---------------------------------------------------------------------------

def test_cen_vs_sap_borrador_end_to_end():
    """Fixtures del mismo periodo (junio): debe haber cruce real."""
    result = run_profile(
        _cen_profile(),
        left_path=CEN_JUNIO,
        right_path=SAP_MUESTRA,
    )
    s = result.summary()
    assert s["matched"] > 0
    assert s["no_cruzados"] == s["solo_left"] + s["solo_right"]
    assert result.kpis["cumplimiento_entregas_pct"] is not None
    assert result.kpis["unidades_pedidas"] > 0
    assert result.kpis["unidades_entregadas"] > 0


def test_cen_grano_agrupado_no_duplica_ordenes():
    result = run_profile(
        _cen_profile(),
        left_path=CEN_JUNIO,
        right_path=SAP_MUESTRA,
    )
    # tras group_by (numero_orden, codigo_item) no debe haber repetidos
    matched_plus_left = result.summary()["matched"] + result.summary()["solo_left"]
    assert matched_plus_left == result.left_meta["num_filas_post_transform"]


# ---------------------------------------------------------------------------
# Proteccion de grano: keys duplicadas sin group_by -> error explicativo
# ---------------------------------------------------------------------------

def test_keys_duplicadas_sin_groupby_lanza_grano_error():
    raw = json.loads(
        (PROFILES_DIR / "cen_vs_sap_v1_borrador.json").read_text(encoding="utf-8")
    )
    raw["left"]["transforms"] = []  # quitar el group_by del CEN
    profile = MatchProfile.model_validate(raw)
    with pytest.raises(GranoNoResueltoError, match="group_by_aggregate"):
        run_profile(profile, left_path=CEN_XLSX, right_path=SAP_MUESTRA)


def test_auto_fix_grano_corrige_group_by_con_columnas_extra():
    """Si el agente mete descriptivas en `by`, el orquestador las corrige."""
    from app.agents.orchestrator import _auto_fix_grano

    raw = json.loads(
        (PROFILES_DIR / "cen_vs_sap_v1_borrador.json").read_text(encoding="utf-8")
    )
    raw["left"]["transforms"][0]["by"] = [
        "numero_orden",
        "codigo_item",
        "descripcion_item",
        "razon_social_comprador",
    ]
    profile = MatchProfile.model_validate(raw)
    with pytest.raises(GranoNoResueltoError):
        run_profile(profile, left_path=CEN_JUNIO, right_path=SAP_MUESTRA)

    fixed, fixes = _auto_fix_grano(profile, CEN_JUNIO, SAP_MUESTRA)
    assert fixes
    assert fixed.left.transforms[-1].by == ["numero_orden", "codigo_item"]
    result = run_profile(fixed, left_path=CEN_JUNIO, right_path=SAP_MUESTRA)
    assert result.summary()["matched"] >= 1


# ---------------------------------------------------------------------------
# Cero perdida con fixture sintetico (gate duro de rubrica Fase 1)
# ---------------------------------------------------------------------------

def _mini_xlsx(tmp_path: Path, name: str, headers: list[str], rows: list[list]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    path = tmp_path / name
    wb.save(path)
    return path


def _mini_profile() -> MatchProfile:
    return MatchProfile.model_validate(
        {
            "profile_id": "mini",
            "version": 1,
            "left": {
                "role": "plan",
                "label": "Plan",
                "loader": {
                    "type": "tabular",
                    "header_row": 1,
                    "columns": [
                        {"name": "k", "source": "K", "dtype": "str"},
                        {"name": "plan", "source": "PLAN", "dtype": "float_clean"},
                    ],
                },
            },
            "right": {
                "role": "real",
                "label": "Real",
                "loader": {
                    "type": "tabular",
                    "header_row": 1,
                    "columns": [
                        {"name": "k", "source": "K", "dtype": "str"},
                        {"name": "real", "source": "REAL", "dtype": "float_clean"},
                    ],
                },
            },
            "join": {"keys": [{"left": "k", "right": "k", "normalizers": ["strip"]}]},
            "computed": [
                {"name": "cumpl", "op": "ratio_pct", "left": "real", "right": "plan", "round": 2}
            ],
            "kpis": [
                {
                    "id": "cumpl_global",
                    "label": "Cumplimiento",
                    "op": "ratio_pct_of_sums",
                    "numerator": "real",
                    "denominator": "plan",
                }
            ],
        }
    )


def test_cero_perdida_sintetico(tmp_path):
    left = _mini_xlsx(
        tmp_path, "plan.xlsx", ["K", "PLAN"],
        [["A", 100], ["B", 200], ["C", 50]],
    )
    right = _mini_xlsx(
        tmp_path, "real.xlsx", ["K", "REAL"],
        [["A", 90], ["B", 210], ["D", 999]],
    )
    result = run_profile(_mini_profile(), left_path=left, right_path=right)
    assert result.summary() == {
        "matched": 2, "solo_left": 1, "solo_right": 1, "no_cruzados": 2,
    }
    assert result.kpis["cumpl_global"] == 100.0
    fila_a = result.matched[result.matched["_k0"] == "A"].iloc[0]
    assert fila_a["cumpl"] == 90.0
    origenes = set(result.no_cruzados["origen"])
    assert origenes == {"plan", "real"}


def test_normalizador_lstrip_zeros_cruza_codigos_con_padding(tmp_path):
    left = _mini_xlsx(tmp_path, "l.xlsx", ["K", "PLAN"], [["00030049", 10]])
    right = _mini_xlsx(tmp_path, "r.xlsx", ["K", "REAL"], [["30049", 8]])
    profile = _mini_profile().model_copy(deep=True)
    profile.join.keys[0].normalizers = [
        *profile.join.keys[0].normalizers,
    ]
    profile.join.keys[0].normalizers.append("lstrip_zeros")
    result = run_profile(profile, left_path=left, right_path=right)
    assert result.summary()["matched"] == 1


def test_grano_consistente_con_normalizadores_no_idempotentes(tmp_path):
    """Regresion del 422 falso: cuando el grano se resuelve con
    group_by_aggregate sobre la llave cruda y el join usa normalizadores no
    idempotentes (lstrip_zeros + digits_only), placeholders como 'SIN DC' y
    celdas vacias colapsaban a llaves distintas tras el group_by pero a la
    MISMA llave tras el join -> GranoNoResueltoError espurio. La
    pre-normalizacion a punto fijo debe evitarlo."""
    left = _mini_xlsx(
        tmp_path, "cen.xlsx", ["ORDEN", "MAT", "PLAN"],
        [["004-0018849", "30049", 10], ["004-0018849", "030049", 5]],
    )
    right = _mini_xlsx(
        tmp_path, "sap.xlsx", ["PEDIDO", "MAT", "REAL"],
        [
            ["004-0018849", "030049", 8],  # cruza con la orden CEN (normalizada)
            ["SIN DC", "30049", 3],        # placeholder -> '' -> '0' al reaplicar
            ["*", "30049", 1],             # placeholder -> '' -> '0'
        ],
    )
    profile = MatchProfile.model_validate(
        {
            "profile_id": "grano_norm",
            "version": 1,
            "left": {
                "role": "cen", "label": "CEN",
                "loader": {
                    "type": "tabular", "header_row": 1,
                    "columns": [
                        {"name": "orden", "source": "ORDEN", "dtype": "str"},
                        {"name": "mat", "source": "MAT", "dtype": "str"},
                        {"name": "plan", "source": "PLAN", "dtype": "float_clean"},
                    ],
                },
                "transforms": [
                    {
                        "op": "group_by_aggregate",
                        "by": ["orden", "mat"],
                        "aggregations": [
                            {"target": "plan", "source": "plan", "fn": "sum"}
                        ],
                    }
                ],
            },
            "right": {
                "role": "sap", "label": "SAP",
                "loader": {
                    "type": "tabular", "header_row": 1,
                    "columns": [
                        {"name": "pedido", "source": "PEDIDO", "dtype": "str"},
                        {"name": "mat", "source": "MAT", "dtype": "str"},
                        {"name": "real", "source": "REAL", "dtype": "float_clean"},
                    ],
                },
                "transforms": [
                    {
                        "op": "group_by_aggregate",
                        "by": ["pedido", "mat"],
                        "aggregations": [
                            {"target": "real", "source": "real", "fn": "sum"}
                        ],
                    }
                ],
            },
            "join": {
                "keys": [
                    {"left": "orden", "right": "pedido",
                     "normalizers": ["strip", "lstrip_zeros", "digits_only"]},
                    {"left": "mat", "right": "mat",
                     "normalizers": ["strip", "lstrip_zeros", "digits_only"]},
                ]
            },
            "kpis": [{"id": "n", "label": "n", "op": "count"}],
        }
    )
    # No debe lanzar GranoNoResueltoError: los placeholders/vacios colapsan a
    # una sola llave '0' de forma consistente entre group_by y join.
    result = run_profile(profile, left_path=left, right_path=right)
    s = result.summary()
    assert s["matched"] == 1  # la orden real cruza
    # los placeholders ('SIN DC', '*') colapsan consistentemente a la misma
    # llave '0' entre group_by y join: 1 sola fila solo_right (sin 422)
    assert s["solo_right"] == 1


def test_unpivot_layout_matriz(tmp_path):
    """Gate Fase 1: layouts matriz (una columna por formato) a formato largo."""
    matriz = _mini_xlsx(
        tmp_path, "matriz.xlsx",
        ["REFERENCIA", "ESTUCHE", "AMARRADO", "VITAFILM"],
        [["ORO", 10, 20, None], ["PLUS", 5, None, 7]],
    )
    profile = _mini_profile().model_copy(deep=True)
    profile_dict = json.loads(profile.to_json())
    profile_dict["left"]["loader"]["columns"] = [
        {"name": "referencia", "source": "REFERENCIA", "dtype": "str"},
        {"name": "estuche", "source": "ESTUCHE", "dtype": "float_clean", "required": False},
        {"name": "amarrado", "source": "AMARRADO", "dtype": "float_clean", "required": False},
        {"name": "vitafilm", "source": "VITAFILM", "dtype": "float_clean", "required": False},
    ]
    profile_dict["left"]["transforms"] = [
        {
            "op": "unpivot",
            "id_vars": ["referencia"],
            "var_name": "formato",
            "value_name": "plan",
        },
        {
            "op": "group_by_aggregate",
            "by": ["referencia"],
            "aggregations": [{"target": "plan", "source": "plan", "fn": "sum"}],
        },
    ]
    profile_dict["join"]["keys"] = [{"left": "referencia", "right": "k"}]
    profile = MatchProfile.model_validate(profile_dict)

    right = _mini_xlsx(tmp_path, "real.xlsx", ["K", "REAL"], [["ORO", 25], ["PLUS", 12]])
    result = run_profile(profile, left_path=matriz, right_path=right)
    assert result.summary()["matched"] == 2
    fila_oro = result.matched[result.matched["_k0"] == "ORO"].iloc[0]
    # float_clean rellena celdas vacias con 0: 10 + 20 + 0
    assert fila_oro["plan"] == 30.0


def test_division_por_cero_da_nan_no_explota(tmp_path):
    left = _mini_xlsx(tmp_path, "l.xlsx", ["K", "PLAN"], [["A", 0]])
    right = _mini_xlsx(tmp_path, "r.xlsx", ["K", "REAL"], [["A", 5]])
    result = run_profile(_mini_profile(), left_path=left, right_path=right)
    import pandas as pd

    assert pd.isna(result.matched.iloc[0]["cumpl"])
    assert result.kpis["cumpl_global"] is None


def test_run_profile_multi_un_par_equivale_a_run_profile():
    """Con un solo archivo por lado, run_profile_multi da el MISMO resultado
    que run_profile (misma union, mismo cruce, misma contabilidad)."""
    profile = _cen_profile()
    single = run_profile(profile, left_path=CEN_JUNIO, right_path=SAP_MUESTRA)
    multi = run_profile_multi(profile, [CEN_JUNIO], [SAP_MUESTRA])
    assert single.summary() == multi.summary()
    assert single.kpis.get("cumplimiento_entregas_pct") == multi.kpis.get(
        "cumplimiento_entregas_pct"
    )


def test_run_profile_multi_consolida_varios_archivos():
    """Unir el mismo par dos veces duplica el volumen de entrada y mantiene la
    garantia de cero perdida (verify_accounting no lanza)."""
    profile = _cen_profile()
    single = run_profile(profile, left_path=CEN_JUNIO, right_path=SAP_MUESTRA)
    doble = run_profile_multi(
        profile, [CEN_JUNIO, CEN_JUNIO], [SAP_MUESTRA, SAP_MUESTRA]
    )
    s1 = single.summary()
    sd = doble.summary()
    # al duplicar la entrada, el total contabilizado crece (no hay perdida)
    assert sd["matched"] + sd["solo_left"] >= s1["matched"] + s1["solo_left"]


def test_auto_fix_join_keys_remapea_columna_clean_inexistente():
    """Si el agente nombra una join key 'X_clean' que no existe, el auto-fix la
    remapea a la columna base real 'X' y arrastra el rename al group_by."""
    from app.agents.orchestrator import _auto_fix_join_keys

    profile = MatchProfile.model_validate(
        {
            "profile_id": "t_keys",
            "version": 1,
            "left": {
                "role": "a",
                "label": "A",
                "loader": {
                    "type": "tabular",
                    "header_row": 1,
                    "columns": [
                        {"name": "codigo_material", "source": "MAT", "dtype": "str"},
                        {"name": "cantidad", "source": "CANT", "dtype": "float_clean"},
                    ],
                },
                "transforms": [
                    {
                        "op": "group_by_aggregate",
                        "by": ["codigo_material_clean"],
                        "aggregations": [
                            {"target": "cantidad", "source": "cantidad", "fn": "sum"}
                        ],
                    }
                ],
            },
            "right": {
                "role": "b",
                "label": "B",
                "loader": {
                    "type": "tabular",
                    "header_row": 1,
                    "columns": [
                        {"name": "material", "source": "M", "dtype": "str"},
                        {"name": "real", "source": "R", "dtype": "float_clean"},
                    ],
                },
                "transforms": [],
            },
            "join": {
                "keys": [
                    {"left": "codigo_material_clean", "right": "material", "normalizers": ["strip"]}
                ],
                "type": "outer",
            },
            "kpis": [{"id": "c", "label": "Cuenta", "op": "count"}],
        }
    )
    fixed, fixes = _auto_fix_join_keys(profile)
    assert fixes, "debio corregir la llave inexistente"
    assert fixed.join.keys[0].left == "codigo_material"
    # el group_by tambien se remapeo al nombre real
    gb = fixed.left.transforms[0]
    assert gb.by == ["codigo_material"]


def test_group_by_sum_tolera_valores_no_numericos():
    """Un group_by con sum sobre una columna que trae basura de texto
    ('CC55', vacios) no debe reventar con 'could not convert string to
    float': los valores no numericos se coercionan a NaN antes de sumar."""
    import pandas as pd

    from app.platform.engine import _apply_transform
    from app.platform.profile import GroupByAggregate

    df = pd.DataFrame(
        {
            "orden": ["A", "A", "B", "B"],
            "cantidad": ["10", "CC55", "5", ""],
            "desc": ["x", "y", "z", "w"],
        }
    )
    transform = GroupByAggregate.model_validate(
        {
            "op": "group_by_aggregate",
            "by": ["orden"],
            "aggregations": [
                {"target": "total", "source": "cantidad", "fn": "sum"},
                {"target": "desc", "source": "desc", "fn": "first"},
            ],
        }
    )
    out = _apply_transform(df, transform, {})
    total = {r["orden"]: r["total"] for _, r in out.iterrows()}
    assert total["A"] == 10.0  # 'CC55' -> NaN, no rompe
    assert total["B"] == 5.0   # '' -> NaN
    # 'first' sobre la columna de texto no se ve afectado por la coercion
    assert set(out["desc"]) == {"x", "z"}
